#!/usr/bin/env python
"""Prepare the ILO **ISCO-08** occupation classification as a labeled benchmark.

This turns two public ILO workbooks into the package's standard inputs:

    <out>/classes.csv     key,description        (flat, one glued string)
    <out>/classes.jsonl   structured taxonomy    (one JSON object per class)
    <out>/items.csv       text,label             (7k job titles -> ISCO codes)

Why this dataset matters: it is a *real official statistical classification*
with a *public* labeled item set, which is rare. The occupation-coding
literature (national statistical offices coding survey write-ins to ISCO/SOC)
almost universally works on confidential microdata, so results are not
reproducible. This one is downloadable, permanently hosted by the ILO, and has
exactly the shape this package targets:

- **436 unit groups**, each with an official title, a definition, a task list,
  and — most usefully for retrieval — an explicit list of example occupations.
- **~7,000 labeled job titles**, naturally imbalanced: median 13 titles per
  class, but a long tail down to 1 and a head up to 113. Class imbalance here
  is a property of the source, not something we induced by subsampling.
- **A second revision of the same taxonomy.** The index carries each title's
  ISCO-88 code alongside its ISCO-08 one, so ``--target isco88`` produces the
  same items labeled into the *previous* revision. That is the taxonomy-revision
  experiment (does a model trained on one revision survive the other, and can
  new classes be absorbed without a retrain) on real data, for free.

Two sources, both from the ILO's ISCO-08 download page
(<https://isco-ilo.netlify.app/en/isco-08/>):

``ISCO-08 EN Structure and definitions.xlsx``
    619 rows covering all four hierarchy levels (10 major groups -> 43 sub-major
    -> 130 minor -> 436 unit groups), with ``Definition``, ``Tasks include``,
    ``Included occupations`` and ``Excluded occupations`` as separate columns.
    This becomes the class definitions.

``ISCO-08 -88 EN Index.xlsx``
    7,018 rows of ``English title`` -> ``ISCO-08`` + ``ISCO-88`` code. This is
    the ILO's own coding index — the alphabetical list a human coder would look
    a write-in up in. It becomes the labeled items.

Both ``classes.csv`` and ``classes.jsonl`` are written every run. The CSV glues
title/definition/inclusions into one description (what the plain
description-similarity signal indexes); the JSONL keeps the fields apart so each
can be indexed as its own short retrieval document. The JSONL is the better
input — pass it to ``--classes`` directly.

Usage:
    python prepare.py [--out build] [--level 4] [--target isco08]
    python prepare.py --level 3            # minor groups (130 classes)
    python prepare.py --major-groups 2,3   # professionals + technicians only
    python prepare.py --target isco88      # same items, previous revision
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import urllib.request

BASE = "https://www.ilo.org/ilostat-files/ISCO/newdocs-08-2021/ISCO-08/"
STRUCTURE_URL = BASE + "ISCO-08%20EN%20Structure%20and%20definitions.xlsx"
INDEX_URL = BASE + "ISCO-08%20-88%20EN%20Index.xlsx"

STRUCTURE_FILE = "ISCO-08-structure-and-definitions.xlsx"
INDEX_FILE = "ISCO-08-88-index.xlsx"

# The description glue caps the inclusion list, for the same reason the COICOP
# preparer does: a long inclusion list dilutes the definition it is glued to and
# BM25 length-normalization penalizes the resulting document twice over. The
# structured JSONL keeps the full list.
MAX_INCLUSIONS_CHARS = 400

# "Tasks include - (a) presiding over ...; (b) ..." -> strip the boilerplate lead.
_TASKS_LEAD = re.compile(r"^\s*Tasks\s+include\s*[-–:]?\s*", re.IGNORECASE)
# "Examples of the occupations classified here: - City councillor - Mayor - ..."
_EXAMPLES_LEAD = re.compile(
    r"^\s*(Examples of the occupations classified here|Some related occupations "
    r"classified elsewhere)\s*:?\s*",
    re.IGNORECASE,
)
# The workbooks use " - " as the bullet separator inside a single cell.
_BULLET = re.compile(r"\s+-\s+")


def _download(url: str, dest: str) -> None:
    if os.path.isfile(dest):
        return
    print(f"downloading {os.path.basename(dest)} ...")
    os.makedirs(os.path.dirname(dest) or ".", exist_ok=True)
    # ilo.org 403s the default `Python-urllib/x.y` User-Agent, so send a browser
    # one. Read the body fully before writing: urlretrieve's Content-Length check
    # trips behind chunking proxies, and a half-written xlsx is worse than none.
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=180) as resp:
        body = resp.read()
    with open(dest, "wb") as fh:
        fh.write(body)


def _clean(value: object) -> str:
    """Workbook cell -> a single-spaced string ('' for blank/NaN)."""
    if value is None:
        return ""
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return ""
    return re.sub(r"\s+", " ", text)


def _code(value: object, width: int) -> str:
    """A cell holding an ISCO code -> a zero-padded string of ``width`` digits.

    Major group 0 (armed forces) is why this exists: its unit groups are ``0110``,
    ``0210``, ``0310``, and openpyxl hands back the *integer* 110 for a cell the
    workbook stores as a number. Left as-is, those three classes silently fail to
    join against the structure file, which stores them as text with the zero
    intact. Non-numeric codes are passed through untouched.
    """
    text = _clean(value)
    if not text or not text.isdigit():
        return text
    return text.zfill(width)


def _bullets(value: object, lead: re.Pattern[str]) -> list[str]:
    """Split one bulleted cell into its items, dropping the boilerplate lead."""
    text = _clean(value)
    if not text:
        return []
    text = lead.sub("", text)
    parts = [p.strip(" -–;.") for p in _BULLET.split(text)]
    return [p for p in parts if p]


def _read_rows(path: str, sheet: int = 0) -> tuple[list[str], list[list[object]]]:
    """(header, rows) from the first worksheet, via openpyxl only.

    Deliberately not pandas: this script is the on-ramp to the example, and the
    core package install is pandas-light. openpyxl is the one extra a user needs.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - user-facing install hint
        raise SystemExit(
            "error: this script needs openpyxl to read the ILO workbooks.\n  pip install openpyxl"
        )
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[sheet]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        raise SystemExit(f"error: {path!r} sheet {sheet} is empty")
    header = [_clean(c) for c in rows[0]]
    return header, [list(r) for r in rows[1:]]


def _column(header: list[str], *candidates: str) -> int:
    """Index of the first matching column, case/space-insensitively."""
    normalized = [h.lower().replace(" ", "") for h in header]
    for candidate in candidates:
        target = candidate.lower().replace(" ", "")
        if target in normalized:
            return normalized.index(target)
    raise SystemExit(
        f"error: none of {candidates} found in worksheet columns {header!r}. "
        "The ILO may have changed the workbook layout."
    )


def load_structure(path: str) -> dict[str, dict[str, object]]:
    """Every ISCO-08 group, keyed by code, with its level and text fields."""
    header, rows = _read_rows(path)
    i_level = _column(header, "Level")
    i_code = _column(header, "ISCO 08 Code", "ISCO-08 Code", "Code")
    i_title = _column(header, "Title EN", "Title")
    i_def = _column(header, "Definition")
    i_tasks = _column(header, "Tasks include", "Tasks")
    i_incl = _column(header, "Included occupations", "Included")
    i_excl = _column(header, "Excluded occupations", "Excluded")

    groups: dict[str, dict[str, object]] = {}
    for row in rows:
        level = _clean(row[i_level])
        if not level.isdigit():
            continue
        # A group's code has exactly as many digits as its level.
        code = _code(row[i_code], int(level))
        if not code:
            continue
        groups[code] = {
            "level": int(level),
            "title": _clean(row[i_title]),
            "definition": _clean(row[i_def]),
            "tasks": _bullets(row[i_tasks], _TASKS_LEAD),
            "inclusions": _bullets(row[i_incl], _EXAMPLES_LEAD),
            "exclusions": _bullets(row[i_excl], _EXAMPLES_LEAD),
        }
    if not groups:
        raise SystemExit(f"error: no ISCO groups parsed from {path!r}")
    return groups


def load_index(path: str) -> list[tuple[str, str, str]]:
    """(title, isco08, isco88) per row of the ILO coding index."""
    header, rows = _read_rows(path)
    i_08 = _column(header, "ISCO-08", "ISCO 08")
    i_88 = _column(header, "ISCO-88", "ISCO 88")
    i_title = _column(header, "English title", "Title EN", "Title")

    out: list[tuple[str, str, str]] = []
    for row in rows:
        title = _clean(row[i_title])
        code08 = _code(row[i_08], 4)
        code88 = _code(row[i_88], 4)
        if not title or not code08:
            continue
        out.append((title, code08, code88))
    return out


def parent_path(code: str, groups: dict[str, dict[str, object]]) -> list[str]:
    """Titles of the ancestors of ``code``, outermost first.

    ISCO codes are positional — ``2511`` sits under ``251`` under ``25`` under
    ``2`` — so the hierarchy is prefixes of the code, no lookup table needed.
    """
    path: list[str] = []
    for length in range(1, len(code)):
        ancestor = groups.get(code[:length])
        if ancestor:
            path.append(str(ancestor["title"]))
    return path


def build_description(code: str, group: dict[str, object]) -> str:
    """The single glued string the flat ``classes.csv`` carries.

    Ordering is deliberate: title first (the strongest short signal), then the
    definition, then the concrete example occupations. That example list is the
    part a short write-in like "lathe operator" actually echoes, which is why it
    is included at all rather than left to the JSONL.
    """
    parts: list[str] = [str(group["title"])]
    definition = str(group["definition"])
    if definition:
        parts.append(definition)
    inclusions = list(group["inclusions"])  # type: ignore[arg-type]
    if inclusions:
        joined = "; ".join(inclusions)
        if len(joined) > MAX_INCLUSIONS_CHARS:
            joined = joined[:MAX_INCLUSIONS_CHARS].rsplit(";", 1)[0]
        parts.append(f"Examples: {joined}.")
    return " ".join(p.strip() for p in parts if p.strip())


def write_classes(
    out_dir: str,
    codes: list[str],
    groups: dict[str, dict[str, object]],
) -> None:
    csv_path = os.path.join(out_dir, "classes.csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["key", "description"])
        for code in codes:
            writer.writerow([code, build_description(code, groups[code])])

    jsonl_path = os.path.join(out_dir, "classes.jsonl")
    with open(jsonl_path, "w", encoding="utf-8") as fh:
        for code in codes:
            group = groups[code]
            definition = str(group["definition"])
            tasks = list(group["tasks"])  # type: ignore[arg-type]
            if tasks:
                # The task list is a second definitional paragraph, not an
                # example list — it belongs with the definition, not in
                # `examples` (which the retrieval views treat as short strings).
                definition = f"{definition} Tasks include: {'; '.join(tasks)}."
            record = {
                "key": code,
                "description": build_description(code, group),
                "title": str(group["title"]),
                "definition": definition.strip(),
                "inclusions": list(group["inclusions"]),  # type: ignore[arg-type]
                "exclusions": list(group["exclusions"]),  # type: ignore[arg-type]
                "parent_path": parent_path(code, groups),
            }
            fh.write(json.dumps(record, ensure_ascii=False) + "\n")

    print(f"wrote {csv_path} and {jsonl_path} ({len(codes)} classes)")


def write_items(
    out_dir: str,
    index: list[tuple[str, str, str]],
    valid: dict[str, str],
    known_units: set[str],
    target: str,
    min_examples: int = 0,
) -> None:
    """``items.csv`` (text,label), with every row's label truncated to the
    requested level and checked against the class list.

    Rows whose code is not a real unit group are dropped and counted. The ILO
    index contains a small number of these (at the time of writing, exactly one
    row carries the code ``1`` and the title ``n`` — a data-entry slip). Silently
    keeping them would create a phantom class with one nonsense example.
    """
    column = 1 if target == "isco08" else 2
    rows: list[tuple[str, str]] = []
    dropped_bad_code = 0
    dropped_out_of_scope = 0
    for record in index:
        title = record[0]
        code = record[column]
        if not code or code not in known_units:
            dropped_bad_code += 1
            continue
        label = valid.get(code)
        if label is None:
            # A real unit group, but outside the requested level / major-group
            # subset. Counted separately so the two reasons stay readable.
            dropped_out_of_scope += 1
            continue
        rows.append((title, label))

    if min_examples > 1:
        # Drop *items*, never classes: the class stays in classes.csv with its
        # official description, so it remains a live candidate the model can
        # still reach through the description signal alone. That is the honest
        # representation of a real taxonomy — every code is codeable, but not
        # every code has training data — and it is what makes `--folds k`
        # feasible, since StratifiedKFold needs k examples of every *labeled*
        # class. See the README for the trade-off this encodes.
        sizes: dict[str, int] = {}
        for _, label in rows:
            sizes[label] = sizes.get(label, 0) + 1
        thin = {label for label, n in sizes.items() if n < min_examples}
        if thin:
            before = len(rows)
            rows = [r for r in rows if r[1] not in thin]
            print(
                f"  --min-examples {min_examples}: dropped {before - len(rows)} items "
                f"across {len(thin)} class(es) with too few titles; those classes "
                f"remain in classes.csv with no labeled examples"
            )

    path = os.path.join(out_dir, "items.csv")
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["text", "label"])
        writer.writerows(rows)

    print(
        f"wrote {path} ({len(rows)} items; dropped {dropped_bad_code} malformed, "
        f"{dropped_out_of_scope} out of scope)"
    )
    counts: dict[str, int] = {}
    for _, label in rows:
        counts[label] = counts.get(label, 0) + 1
    if counts:
        sizes = sorted(counts.values())
        median = sizes[len(sizes) // 2]
        print(
            f"  {len(counts)} classes represented; items per class: "
            f"min {sizes[0]}, median {median}, max {sizes[-1]}"
        )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", default="build", help="output directory (default: build)")
    parser.add_argument(
        "--level",
        type=int,
        default=4,
        choices=(1, 2, 3, 4),
        help="hierarchy level to classify into: 1 major (10 classes), 2 sub-major "
        "(43), 3 minor (130), 4 unit groups (436, the default)",
    )
    parser.add_argument(
        "--major-groups",
        default="",
        help="comma-separated major groups to keep, e.g. '2,3' (default: all)",
    )
    parser.add_argument(
        "--min-examples",
        type=int,
        default=0,
        help="drop items belonging to classes with fewer than this many labeled "
        "titles (default: 0, keep everything). The classes themselves are kept "
        "in classes.csv, so they stay predictable from their description alone. "
        "Set this to your --folds value: StratifiedKFold needs k examples of "
        "every labeled class, and at level 4 eight classes have fewer than 3.",
    )
    parser.add_argument(
        "--target",
        default="isco08",
        choices=("isco08", "isco88"),
        help="which revision's codes to use as labels (default: isco08). "
        "'isco88' labels the same items into the previous revision — the "
        "taxonomy-revision experiment. Class definitions are always ISCO-08, so "
        "use isco88 for item-side comparisons only.",
    )
    args = parser.parse_args(argv)

    os.makedirs(args.out, exist_ok=True)
    structure_path = os.path.join(args.out, STRUCTURE_FILE)
    index_path = os.path.join(args.out, INDEX_FILE)
    _download(STRUCTURE_URL, structure_path)
    _download(INDEX_URL, index_path)

    groups = load_structure(structure_path)
    index = load_index(index_path)

    keep_majors = {m.strip() for m in args.major_groups.split(",") if m.strip()}
    codes = sorted(
        code
        for code, group in groups.items()
        if group["level"] == args.level and (not keep_majors or code[0] in keep_majors)
    )
    if not codes:
        raise SystemExit(
            f"error: no groups at level {args.level} for major groups "
            f"{sorted(keep_majors) or 'all'}"
        )

    if args.target == "isco08":
        known_units = {code for code, group in groups.items() if group["level"] == 4}
        in_scope = set(codes)
    else:
        # ISCO-88 codes have no definitions in this workbook, so classes stay
        # ISCO-08 and only the item labels change. Say so loudly rather than
        # writing a classes/items pair that cannot be trained on as-is.
        print(
            "note: --target isco88 relabels items into ISCO-88 while classes.csv "
            "stays ISCO-08. Use it for revision-shift analysis, not as a direct "
            "train/classes pair."
        )
        known_units = {record[2] for record in index if len(record[2]) == 4}
        in_scope = {
            code[: args.level] if args.level < 4 else code
            for code in known_units
            if not keep_majors or code[0] in keep_majors
        }

    # Map every known unit group to the label it collapses to at this level.
    valid = {
        code: label
        for code, label in ((c, c[: args.level] if args.level < 4 else c) for c in known_units)
        if label in in_scope
    }

    write_classes(args.out, codes, groups)
    write_items(args.out, index, valid, known_units, args.target, args.min_examples)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
